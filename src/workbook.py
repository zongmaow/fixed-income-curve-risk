"""Excel workbook: bond risk as values, portfolio roll-up as formulas."""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from .bonds import INSTRUMENT_IDS
from .scenarios import SCENARIO_ORDER


BOOK_ORDER = (
    "current",
    "benchmark",
    "candidate_10pct_s05y",
    "candidate_expanded_s05y",
    "control_2y20y",
)
BOOK_LABELS = {
    "current": "Current",
    "benchmark": "Benchmark",
    "candidate_10pct_s05y": "Candidate 10%",
    "candidate_expanded_s05y": "Expanded S05Y",
    "control_2y20y": "Control 2y/20y",
}
APPLY_COST = {
    "current": 0,
    "benchmark": 0,
    "candidate_10pct_s05y": 1,
    "candidate_expanded_s05y": 1,
    "control_2y20y": 0,
}


def _header(cell):
    cell.font = Font(bold=True)
    cell.alignment = Alignment(wrap_text=True, vertical="center")


def build_workbook(path, instruments: dict, books: dict, bond_pnl_per_100: dict, five_year_detail: list[dict], breakeven_bp: float):
    wb = Workbook()

    inputs = wb.active
    inputs.title = "Inputs"
    inputs["A1"] = "Case inputs"
    inputs["A1"].font = Font(bold=True, size=14)
    inputs["A3"] = "Initial NAV"
    inputs["B3"] = 100_000_000
    inputs["A4"] = "DV01 target per bp"
    inputs["B4"] = 50_000
    inputs["A5"] = "Relative loss budget"
    inputs["B5"] = 1_000_000
    inputs["A6"] = "One-sided transaction cost (bp)"
    inputs["B6"] = 2
    inputs["A7"] = "Routine turnover cap"
    inputs["B7"] = 0.10
    inputs["A9"] = "Edit B6, the weight block, and the apply-cost flags. Bond prices and scenario returns are values from the Python revaluation. NAV, DV01, turnover, cost, and relative P&L are formulas."
    inputs["A9"].alignment = Alignment(wrap_text=True)
    inputs.column_dimensions["A"].width = 42
    inputs.column_dimensions["B"].width = 18
    inputs["B3"].number_format = '#,##0'
    inputs["B4"].number_format = '#,##0'
    inputs["B5"].number_format = '#,##0'
    inputs["B6"].number_format = '0.00'
    inputs["B7"].number_format = '0.00%'
    for row in range(3, 8):
        _header(inputs.cell(row, 1))

    wb.defined_names.add(DefinedName(name="NAV", attr_text="Inputs!$B$3"))
    wb.defined_names.add(DefinedName(name="DV01Target", attr_text="Inputs!$B$4"))
    wb.defined_names.add(DefinedName(name="Budget", attr_text="Inputs!$B$5"))
    wb.defined_names.add(DefinedName(name="CostBp", attr_text="Inputs!$B$6"))

    bonds = wb.create_sheet("Bonds")
    bond_headers = [
        "Instrument",
        "Coupon",
        "Maturity",
        "Dirty per 100",
        "DV01 per 100",
        "DV01 per dollar",
    ] + list(SCENARIO_ORDER)
    for col, header in enumerate(bond_headers, 1):
        cell = bonds.cell(1, col, header)
        _header(cell)
    for i, name in enumerate(INSTRUMENT_IDS):
        inst = instruments[name]
        row = 2 + i
        bonds.cell(row, 1, name)
        bonds.cell(row, 2, inst["coupon_percent"] / 100.0).number_format = '0.00%'
        bonds.cell(row, 3, inst["maturity"])
        bonds.cell(row, 4, inst["dirty"]).number_format = '0.000000'
        bonds.cell(row, 5, inst["dv01_per_100"]).number_format = '0.00000000'
        bonds.cell(row, 6, f"=E{row}/D{row}").number_format = '0.00000000'
        for s_index, scenario in enumerate(SCENARIO_ORDER):
            ret = bond_pnl_per_100[name][scenario] / inst["dirty"]
            cell = bonds.cell(row, 7 + s_index, ret)
            cell.number_format = '0.00000000'
    bonds.auto_filter.ref = f"A1:L{1+len(INSTRUMENT_IDS)}"
    bonds.freeze_panes = "B2"
    for col in range(1, 13):
        bonds.column_dimensions[get_column_letter(col)].width = 22 if col > 6 else 18

    weights = wb.create_sheet("Weights")
    weights["A1"] = "Instrument"
    _header(weights["A1"])
    for col, book in enumerate(BOOK_ORDER, 2):
        cell = weights.cell(1, col, BOOK_LABELS[book])
        _header(cell)
    for i, name in enumerate(INSTRUMENT_IDS):
        row = 2 + i
        weights.cell(row, 1, name)
        for col, book in enumerate(BOOK_ORDER, 2):
            cell = weights.cell(row, col, books[book]["weights"][name])
            cell.number_format = '0.000000%'
    weights["A11"] = "Apply one-sided cost"
    _header(weights["A11"])
    weights["A12"] = "1 charges buys and sells against this column. Current, benchmark, and control stay at 0."
    for col, book in enumerate(BOOK_ORDER, 2):
        weights.cell(11, col, APPLY_COST[book])
    weights.column_dimensions["A"].width = 28
    for col in range(2, 7):
        weights.column_dimensions[get_column_letter(col)].width = 20

    port = wb.create_sheet("Portfolio")
    port["A1"] = "Item"
    _header(port["A1"])
    for col, book in enumerate(BOOK_ORDER, 2):
        cell = port.cell(1, col, BOOK_LABELS[book])
        _header(cell)
    port["A2"] = "Apply cost"
    for col in range(2, 7):
        letter = get_column_letter(col)
        port.cell(2, col, f"=Weights!{letter}11")

    port["A4"] = "Market value"
    _header(port["A4"])
    for i in range(8):
        row = 5 + i
        port.cell(row, 1, f"=Weights!A{2+i}")
        for col in range(2, 7):
            letter = get_column_letter(col)
            cell = port.cell(row, col, f"=Weights!{letter}{2+i}*NAV")
            cell.number_format = '#,##0.00'
    port["A13"] = "Total market value"
    port["A14"] = "Sum of weights"
    port["A15"] = "Parallel DV01"
    port["A16"] = "DV01 minus target"
    port["A17"] = "DV01 inside ±2%"
    for col in range(2, 7):
        letter = get_column_letter(col)
        port.cell(13, col, f"=SUM({letter}5:{letter}12)").number_format = '#,##0.00'
        port.cell(14, col, f"=SUM(Weights!{letter}2:{letter}9)").number_format = '0.000000%'
        port.cell(15, col, f"=SUMPRODUCT({letter}5:{letter}12,Bonds!$F$2:$F$9)").number_format = '#,##0.0000'
        port.cell(16, col, f"={letter}15-DV01Target").number_format = '#,##0.0000'
        port.cell(17, col, f'=IF(ABS({letter}15/DV01Target-1)<=0.02,"inside","outside")')
    for row in (13, 14, 15, 16, 17):
        _header(port.cell(row, 1))

    port["A19"] = "Gross traded vs current (buys+sells)"
    port["A20"] = "Buys vs current"
    port["A21"] = "Sells vs current"
    port["A22"] = "Turnover vs current"
    port["A23"] = "Transaction cost"
    for col in range(2, 7):
        letter = get_column_letter(col)
        # Current column is B. Deltas against $B.
        port.cell(19, col, f"=SUMPRODUCT(ABS({letter}5:{letter}12-$B$5:$B$12))").number_format = '#,##0.00'
        port.cell(
            20,
            col,
            f"=SUMPRODUCT(({letter}5:{letter}12-$B$5:$B$12>0)*({letter}5:{letter}12-$B$5:$B$12))",
        ).number_format = '#,##0.00'
        port.cell(21, col, f"={letter}19-{letter}20").number_format = '#,##0.00'
        port.cell(22, col, f"=0.5*{letter}19/NAV").number_format = '0.0000%'
        port.cell(23, col, f"={letter}2*(CostBp/10000)*{letter}19").number_format = '#,##0.00'
    for row in range(19, 24):
        _header(port.cell(row, 1))

    # Scenarios. Benchmark column is C (3).
    start = 25
    port.cell(start, 1, "Scenario full-revaluation P&L, relative P&L after cost, and budget status")
    _header(port.cell(start, 1))
    for s_index, scenario in enumerate(SCENARIO_ORDER):
        base = start + 1 + s_index * 3
        ret_col = get_column_letter(7 + s_index)  # G is first scenario on Bonds
        port.cell(base, 1, f"{scenario} P&L after cost")
        port.cell(base + 1, 1, f"{scenario} relative to benchmark")
        port.cell(base + 2, 1, f"{scenario} budget")
        for col in range(2, 7):
            letter = get_column_letter(col)
            pnl = port.cell(base, col, f"=SUMPRODUCT({letter}5:{letter}12,Bonds!${ret_col}$2:${ret_col}$9)-{letter}23")
            pnl.number_format = '#,##0.00'
            rel = port.cell(base + 1, col, f"={letter}{base}-C{base}")
            rel.number_format = '#,##0.00'
            status = port.cell(base + 2, col, f'=IF({letter}{base+1}>=-Budget,"inside","outside")')
        for offset in range(3):
            _header(port.cell(base + offset, 1))

    red = PatternFill("solid", fgColor="F4C7C3")
    last_status = start + 1 + (len(SCENARIO_ORDER) - 1) * 3 + 2
    port.conditional_formatting.add(
        f"B{start+1}:F{last_status}",
        CellIsRule(operator="equal", formula=['"outside"'], fill=red),
    )

    port.column_dimensions["A"].width = 62
    for col in range(2, 7):
        port.column_dimensions[get_column_letter(col)].width = 20
    port.row_dimensions[1].height = 22
    port.freeze_panes = "B5"

    cf = wb.create_sheet("FiveYearCashFlows")
    cf["A1"] = "S05Y cash-flow discount check"
    cf["A1"].font = Font(bold=True, size=14)
    cf["A2"] = "Present value uses Excel EXP on the continuous zero. T and the zero are values from the valuation-date curve. Sum of PV should match the S05Y dirty price on Bonds."
    cf["A2"].alignment = Alignment(wrap_text=True)
    headers = ["Pay date", "T (ACT/365F)", "Cash flow per 100", "Continuous zero (decimal)", "Discount factor", "Present value"]
    for col, header in enumerate(headers, 1):
        cell = cf.cell(4, col, header)
        _header(cell)
    for i, row_data in enumerate(five_year_detail):
        excel_row = 5 + i
        cf.cell(excel_row, 1, row_data["pay_date"].isoformat())
        cf.cell(excel_row, 2, row_data["T"]).number_format = '0.00000000'
        cf.cell(excel_row, 3, row_data["cf"]).number_format = '0.000000'
        cf.cell(excel_row, 4, row_data["zero_decimal"]).number_format = '0.00000000'
        cf.cell(excel_row, 5, f"=EXP(-D{excel_row}*B{excel_row})").number_format = '0.00000000'
        cf.cell(excel_row, 6, f"=C{excel_row}*E{excel_row}").number_format = '0.00000000'
    last = 4 + len(five_year_detail)
    sum_row = last + 2
    cf.cell(sum_row, 5, "Sum of present values")
    _header(cf.cell(sum_row, 5))
    cf.cell(sum_row, 6, f"=SUM(F5:F{last})").number_format = '0.00000000'
    cf.cell(sum_row + 1, 5, "Dirty price from Bonds (S05Y)")
    cf.cell(sum_row + 1, 6, "=INDEX(Bonds!D:D,MATCH(\"S05Y\",Bonds!A:A,0))").number_format = '0.00000000'
    cf.cell(sum_row + 2, 5, "Difference")
    cf.cell(sum_row + 2, 6, f"=F{sum_row}-F{sum_row+1}").number_format = '0.00000000'
    for col in range(1, 7):
        cf.column_dimensions[get_column_letter(col)].width = 32
    cf.row_dimensions[2].height = 32

    notes = wb.create_sheet("Notes")
    notes["A1"] = "How to read the formulas"
    notes["A1"].font = Font(bold=True, size=14)
    lines = [
        "Market value is weight times Inputs!B3. Parallel DV01 is the sum of market value times DV01 per dollar.",
        "Turnover versus the current book is 0.5 * sum of absolute market-value changes / NAV. Gross traded value is buys plus sells.",
        "Cost is the one-sided bp rate times gross traded value, and only for columns with apply-cost = 1.",
        "Scenario P&L is the sum of market value times the bond's full-revaluation return, minus that cost.",
        "Relative P&L subtracts the benchmark column. Inside the budget means relative P&L is not worse than -Budget.",
        "The steepener is -50 bp at and before 2y, +100 bp at and beyond 10y, linear in maturity between those nodes. The flattener swaps the signs. Neither is a bear steepener or a bull flattener.",
        f"At the time this file was written, the expanded S05Y book's tightest scenario had a one-sided cost break-even of {breakeven_bp:.4f} bp. The live formula below recomputes that break-even from the Portfolio sheet. A one-sided cost above it exhausts the cushion. 4.2 bp is the design note's illustration; it exhausts the cushion only if the live break-even is below 4.2.",
    ]
    for i, line in enumerate(lines):
        notes.cell(3 + i, 1, line)
        notes.cell(3 + i, 1).alignment = Alignment(wrap_text=True)
        notes.row_dimensions[3 + i].height = 36
    # Live break-even. Expanded book is column E. Relative rows are start+1+3*s+1 = start+2+3*s
    rel_rows = [start + 2 + 3 * s for s in range(len(SCENARIO_ORDER))]
    rel_cells = ",".join(f"Portfolio!E{row}" for row in rel_rows)
    notes["A12"] = "Live cushion on expanded book (tightest relative P&L + budget)"
    notes["B12"] = f"=MIN({rel_cells})+Budget"
    notes["B12"].number_format = '#,##0.00'
    notes["A13"] = "Live break-even one-sided cost (bp)"
    notes["B13"] = "=CostBp+B12/Portfolio!E19*10000"
    notes["B13"].number_format = '0.0000'
    notes["A14"] = "Would 4.2 bp one-sided exhaust the expanded cushion?"
    notes["B14"] = '=IF(4.2>B13,"yes","no")'
    for row in range(12, 15):
        _header(notes.cell(row, 1))
    notes.column_dimensions["A"].width = 88
    notes.column_dimensions["B"].width = 22

    wb.save(path)
