"""Regression tests for the curve-risk case.

Run from the repository root:

    python3 -m unittest tests/test_case.py
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from src.bonds import (  # noqa: E402
    INSTRUMENT_IDS,
    build_instruments,
    dv01_per_100,
    key_rate_weights,
)
from src.curve import load_curve_parameters, zero_decimal  # noqa: E402
from src.decision import BOOK_ROLES, decide  # noqa: E402
from src.portfolio import (  # noqa: E402
    BUY_CHOICES,
    DV01_TARGET,
    NAV,
    RELATIVE_BUDGET,
    aggregate_clearing,
    book_pnl,
    candidate_stress_cushion,
    candidate_weights,
    normalize_weights,
    portfolio_dv01,
    scenario_return,
    scenario_turnover_bound,
    solve_benchmark,
    solve_two_bond,
    transaction_cost,
    turnover_and_gross,
)
from src.scenarios import SCENARIO_ORDER, build_scenarios, historical_dz  # noqa: E402
from reproduce_case import clearing_analysis, lowest_clearing_buy  # noqa: E402
from src.bonds import full_reval_pnl_per_100  # noqa: E402

DESIGN_PRICES = {
    "S06M": 99.582157,
    "S02Y": 99.194643,
    "S03Y": 99.065240,
    "S05Y": 99.087613,
    "S07Y": 99.566682,
    "S10Y": 99.582423,
    "S20Y": 98.181553,
    "S30Y": 97.092721,
}
JUL_OCT_BP = (0.90, 26.18, 77.43, 108.10, 119.44, 93.82)
MARCH_BP = (-37.57, -95.13, -62.51, -41.39, -21.86, -13.14)


class CaseTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.case = json.loads((ROOT / "outputs" / "case_results.json").read_text())
        curves = load_curve_parameters(ROOT / "data" / "curve_parameters_2022_2023.csv")
        cls.curves = curves
        cls.instruments = build_instruments(curves["2023-07-31"])

    def test_frozen_sample(self):
        self.assertEqual(len(self.curves), 499)
        manifest = json.loads((ROOT / "data" / "source_manifest.json").read_text())
        self.assertEqual(manifest["filter"]["dated_rows"], 520)
        self.assertEqual(manifest["filter"]["dropped_missing_beta_or_tau"], 21)
        self.assertEqual(len(manifest["dropped_dates"]), 21)

    def test_prices_within_1e_6(self):
        for name, expected in DESIGN_PRICES.items():
            dirty = self.instruments[name]["dirty"]
            self.assertLess(abs(dirty - expected), 1e-6, msg=f"{name} {dirty}")

    def test_clean_equals_dirty(self):
        for name in INSTRUMENT_IDS:
            inst = self.instruments[name]
            self.assertEqual(inst["accrued"], 0.0)
            self.assertEqual(inst["clean"], inst["dirty"])

    def test_nav_and_dv01(self):
        for book, payload in self.case["books"].items():
            weights = payload["weights"]
            self.assertAlmostEqual(sum(weights.values()), 1.0, places=10, msg=book)
            self.assertTrue(all(weight >= -1e-12 for weight in weights.values()), book)
            normalized = normalize_weights(weights, INSTRUMENT_IDS)
            dv01 = portfolio_dv01(normalized, self.instruments)
            self.assertAlmostEqual(dv01, DV01_TARGET, delta=1e-6, msg=book)
            self.assertAlmostEqual(sum(payload["market_value"].values()), NAV, delta=1e-4, msg=book)

    def test_control_weights_are_present(self):
        weights = self.case["books"]["control_2y20y"]["weights"]
        self.assertGreater(weights["S02Y"], 0.5)
        self.assertGreater(weights["S20Y"], 0.1)
        self.assertEqual(weights["S30Y"], 0.0)
        self.assertEqual(weights["S06M"], 0.0)

    def test_node_shocks_round_to_published_checks(self):
        checks = self.case["node_shock_checks"]
        self.assertEqual(checks["hist_2023_07_31_to_2023_10_19"]["rounded_bp"], list(JUL_OCT_BP))
        self.assertEqual(checks["hist_2023_03_08_to_2023_03_13"]["rounded_bp"], list(MARCH_BP))

    def test_ten_percent_book_does_not_repair(self):
        book = self.case["books"]["candidate_10pct_s05y"]
        self.assertAlmostEqual(book["turnover_vs_current"], 0.10, places=8)
        self.assertAlmostEqual(book["weights"]["S05Y"], 0.10, places=10)
        failed = [name for name, row in book["scenarios"].items() if not row["within_budget"]]
        self.assertTrue(failed)

    def test_expanded_book_clears_with_cushion(self):
        book = self.case["books"]["candidate_expanded_s05y"]
        self.assertTrue(all(row["within_budget"] for row in book["scenarios"].values()))
        expanded = self.case["expanded_s05y"]
        self.assertGreater(expanded["cushion_dollars"], 0.0)
        self.assertLess(expanded["turnover_just"], book["turnover_vs_current"])
        # One point of NAV turnover was added on top of the just-clearing trade.
        self.assertAlmostEqual(book["turnover_vs_current"] - expanded["turnover_just"], 0.01, places=8)

    def test_s05y_is_not_beaten_by_another_feasible_tenor(self):
        decision = self.case["case_decision"]
        self.assertEqual(decision["directions_that_clear_at_lower_turnover"], [])
        self.assertEqual(decision["lowest_feasible_direction"]["buy"], "S05Y")

    def test_historical_dz_matches_independent_curve_difference(self):
        """Fails if historical_dz is stubbed to return zero."""
        pairs = (("2023-07-31", "2023-10-19"), ("2023-03-08", "2023-03-13"))
        scenarios = build_scenarios(self.curves)
        for start, end in pairs:
            shock = historical_dz(self.curves[start], self.curves[end])
            name = f"hist_{start.replace('-', '_')}_to_{end.replace('-', '_')}"
            # names in the case are hist_2023_07_31_to_2023_10_19, which matches this pattern
            for maturity in (0.5, 2.0, 5.0, 10.0, 20.0, 30.0):
                independent = zero_decimal(maturity, self.curves[end]) - zero_decimal(maturity, self.curves[start])
                self.assertAlmostEqual(shock(maturity), independent, places=12)
                self.assertGreater(abs(independent), 1e-6)
                self.assertAlmostEqual(scenarios[name]["dz"](maturity), independent, places=12)

    def test_key_rate_weights_and_kr01_sum(self):
        for instrument in self.instruments.values():
            kr_sum = sum(instrument["kr01_per_100"])
            self.assertLess(abs(kr_sum - instrument["dv01_per_100"]), 1e-6)
            for flow in instrument["flows"]:
                weights = key_rate_weights(flow["T"])
                self.assertTrue(all(weight >= -1e-15 for weight in weights))
                self.assertAlmostEqual(sum(weights), 1.0, places=12)

    def test_dv01_stable_across_half_one_and_two_bp(self):
        """Normalized central DV01 across 0.5, 1, and 2 bp steps.

        Tolerance is 1e-6 per 100 face, absolute span. On these contracts the
        30-year bond is the widest, and it moves by less than 7e-7.
        """
        beta = self.curves["2023-07-31"]
        for instrument in self.instruments.values():
            values = [dv01_per_100(instrument["flows"], beta, step) for step in (0.5, 1.0, 2.0)]
            self.assertLess(max(values) - min(values), 1e-6, msg=instrument["id"])

    def test_candidate_directions_keep_mandate_and_bounds(self):
        current = normalize_weights(solve_two_bond(self.instruments, "S06M", "S30Y"), INSTRUMENT_IDS)
        benchmark = normalize_weights(solve_benchmark(self.instruments), INSTRUMENT_IDS)
        scenarios = build_scenarios(self.curves)
        from src.bonds import full_reval_pnl_per_100
        from src.portfolio import scenario_return

        beta = self.curves["2023-07-31"]
        returns = {}
        for scenario in SCENARIO_ORDER:
            returns[scenario] = {}
            for name, instrument in self.instruments.items():
                per_100 = full_reval_pnl_per_100(instrument, beta, scenarios[scenario]["dz"])
                returns[scenario][name] = scenario_return(instrument, per_100)
        for buy in BUY_CHOICES:
            weights = normalize_weights(candidate_weights(current, self.instruments, buy, 0.10), INSTRUMENT_IDS)
            self.assertTrue(all(weight >= -1e-12 for weight in weights.values()), buy)
            self.assertAlmostEqual(sum(weights.values()), 1.0, places=10)
            self.assertAlmostEqual(portfolio_dv01(weights, self.instruments), DV01_TARGET, delta=1e-6)
            self.assertAlmostEqual(sum(weight * NAV for weight in weights.values()), NAV, delta=1e-4)
            _turnover, gross, _buys, _sells = turnover_and_gross(weights, current)
            self.assertAlmostEqual(transaction_cost(gross), 0.0002 * gross, delta=1e-6)
            self.assertAlmostEqual(gross, 0.20 * NAV, delta=1e-2)
            screen = clearing_analysis(current, benchmark, self.instruments, returns, buy)
            for row in screen["bounds"]:
                if row["kind"] == "flat":
                    self.assertNotIn("turnover_bound", row)
                    continue
                implied = row["relative_at_zero_turnover"] + row["turnover_bound"] * row["slope_dollars_per_unit_turnover"]
                self.assertAlmostEqual(implied, -RELATIVE_BUDGET, delta=1e-4)

    def test_flat_scenario_inside_budget_has_no_turnover_bound(self):
        row = scenario_turnover_bound(0.0, -100.0)
        row["scenario"] = "flat_inside"
        self.assertEqual(row["kind"], "flat")
        self.assertEqual(row["feasible_if"], "already_inside")
        self.assertNotIn("turnover_bound", row)
        summary = aggregate_clearing([row], 1.0)
        self.assertTrue(summary["feasible"])
        self.assertEqual(summary["turnover_just"], 0.0)
        self.assertEqual(summary["blocked_flat_scenarios"], [])

    def test_flat_scenario_outside_budget_is_infeasible(self):
        row = scenario_turnover_bound(0.0, -2_000_000.0)
        row["scenario"] = "flat_outside"
        self.assertEqual(row["feasible_if"], "impossible")
        self.assertNotIn("turnover_bound", row)
        summary = aggregate_clearing([row], 1.0)
        self.assertFalse(summary["feasible"])
        self.assertEqual(summary["blocked_flat_scenarios"], ["flat_outside"])
        self.assertNotIn("turnover_bound", row)

    def test_saved_json_matches_fresh_shocks_and_decisions(self):
        scenarios = build_scenarios(self.curves)
        for name, spec in scenarios.items():
            if not name.startswith("hist_"):
                continue
            fresh = [round(float(value), 2) for value in spec["node_shock_bp"]]
            self.assertEqual(fresh, self.case["node_shock_checks"][name]["rounded_bp"])
        for name, book in self.case["books"].items():
            relatives = [book["scenarios"][scenario]["relative_pnl_after_cost"] for scenario in SCENARIO_ORDER]
            verdict = decide(
                BOOK_ROLES[name],
                book["weights"],
                book["dv01"],
                relatives,
                book["turnover_vs_current"],
            )
            self.assertEqual(verdict["decision"], book["governance"]["decision"])
            self.assertEqual(verdict["decision"], self.case["decisions"][name])

    def test_contribution_dollars_sum_to_book_pnl(self):
        path = ROOT / "outputs" / "scenario_contributions.csv"
        totals = {}
        with path.open(newline="") as handle:
            for row in csv.DictReader(handle):
                key = (row["book"], row["scenario"])
                totals[key] = totals.get(key, 0.0) + float(row["pnl_dollars"])
        for book, payload in self.case["books"].items():
            for scenario, result in payload["scenarios"].items():
                self.assertAlmostEqual(totals[(book, scenario)], result["pnl"], delta=1e-4)

    def test_readme_says_relieves_does_not_repair(self):
        text = (ROOT / "README.md").read_text()
        self.assertIn("relieves, does not repair", text)

    def test_planted_sveny_gap_fails_reproduce(self):
        sveny_path = ROOT / "data" / "published_sveny_event_dates.csv"
        manifest_path = ROOT / "data" / "source_manifest.json"
        original_sveny = sveny_path.read_bytes()
        original_manifest = manifest_path.read_bytes()
        try:
            rows = list(csv.reader(io.StringIO(original_sveny.decode())))
            column = rows[0].index("SVENY10")
            rows[1][column] = f"{float(rows[1][column]) + 1.0:.10f}"
            buffer = io.StringIO()
            csv.writer(buffer, lineterminator="\n").writerows(rows)
            sveny_path.write_text(buffer.getvalue())
            manifest = json.loads(original_manifest)
            manifest["published_sveny_excerpt_sha256"] = hashlib.sha256(sveny_path.read_bytes()).hexdigest()
            manifest_path.write_text(json.dumps(manifest, indent=2) + "\n")
            proc = subprocess.run(
                [sys.executable, "reproduce_case.py"],
                cwd=ROOT,
                capture_output=True,
                text=True,
            )
            self.assertNotEqual(proc.returncode, 0)
            self.assertIn("Published SVENY", proc.stderr + proc.stdout)
        finally:
            sveny_path.write_bytes(original_sveny)
            manifest_path.write_bytes(original_manifest)

    def test_hash_mismatch_tells_user_to_accept_new_source(self):
        sveny_path = ROOT / "data" / "published_sveny_event_dates.csv"
        original = sveny_path.read_bytes()
        try:
            sveny_path.write_bytes(original + b"\n")
            proc = subprocess.run(
                [sys.executable, "reproduce_case.py"],
                cwd=ROOT,
                capture_output=True,
                text=True,
            )
            self.assertNotEqual(proc.returncode, 0)
            self.assertIn("--accept-new-source", proc.stderr + proc.stdout)
        finally:
            sveny_path.write_bytes(original)

    def test_cost_4_2bp_changes_expanded_decision_after_recalc(self):
        if shutil.which("soffice") is None:
            self.skipTest("LibreOffice is not installed")
        from openpyxl import load_workbook
        from src.workbook import cache_workbook_values

        source = ROOT / "fixed_income_case.xlsx"
        cached = load_workbook(source, data_only=True)
        self.assertEqual(cached["Portfolio"].cell(52, 5).value, "needs_approval")
        self.assertEqual(cached["Portfolio"].cell(40, 5).value, "inside")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"
            shutil.copy(source, path)
            workbook = load_workbook(path)
            self.assertIn("Inputs!$B$7", workbook["Portfolio"].cell(51, 5).value)
            workbook["Inputs"]["B6"] = 4.2
            workbook.save(path)
            self.assertTrue(cache_workbook_values(path))
            recalculated = load_workbook(path, data_only=True)
            self.assertEqual(recalculated["Portfolio"].cell(40, 5).value, "outside")
            self.assertEqual(recalculated["Portfolio"].cell(52, 5).value, "still_outside_budget")

    def test_turnover_cap_raise_makes_expanded_executable_after_recalc(self):
        if shutil.which("soffice") is None:
            self.skipTest("LibreOffice is not installed")
        from openpyxl import load_workbook
        from src.workbook import cache_workbook_values

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"
            shutil.copy(ROOT / "fixed_income_case.xlsx", path)
            workbook = load_workbook(path)
            workbook["Inputs"]["B7"] = 0.50
            workbook.save(path)
            self.assertTrue(cache_workbook_values(path))
            recalculated = load_workbook(path, data_only=True)
            self.assertEqual(recalculated["Portfolio"].cell(52, 5).value, "executable_within_authority")
            self.assertEqual(recalculated["Portfolio"].cell(43, 5).value, "inside")



    def test_dv01_mandate_endpoints_include_dust(self):
        from src.decision import DV01_DUST, DV01_TARGET, dv01_within_mandate
        import validate_independently as independent

        self.assertEqual(DV01_DUST, 1e-6)
        inside = (
            49000.0,
            51000.0,
            DV01_TARGET * 1.02 + 1e-9,
            51000.0000000001,
        )
        outside = (48999.0, 51001.0, 52000.0)
        for value in inside:
            self.assertTrue(dv01_within_mandate(value), msg=value)
            verdict = independent.decision_from_predicates("standing", {"a": 1.0}, value, [-1.0], 0.0)
            self.assertTrue(verdict["dv01_within_mandate"], msg=value)
        for value in outside:
            self.assertFalse(dv01_within_mandate(value), msg=value)
            verdict = independent.decision_from_predicates("standing", {"a": 1.0}, value, [-1.0], 0.0)
            self.assertFalse(verdict["dv01_within_mandate"], msg=value)

    def test_workbook_dv01_dust_endpoint_is_inside_after_recalc(self):
        if shutil.which("soffice") is None:
            self.skipTest("LibreOffice is not installed")
        from openpyxl import load_workbook
        from src.workbook import cache_workbook_values

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"
            shutil.copy(ROOT / "fixed_income_case.xlsx", path)
            workbook = load_workbook(path)
            formula = workbook["Portfolio"].cell(17, 2).value
            mandate = workbook["Portfolio"].cell(49, 2).value
            self.assertIn("DV01Target*0.02+0.000001", formula)
            self.assertIn("DV01Target*0.02+0.000001", mandate)
            self.assertNotIn("/DV01Target-1", formula)
            self.assertNotIn("/DV01Target-1", mandate)
            workbook["Portfolio"]["B15"] = "=51000.0000000001"
            workbook["Inputs"]["B3"] = 102_000_000
            workbook.save(path)
            self.assertTrue(cache_workbook_values(path))
            recalculated = load_workbook(path, data_only=True)
            self.assertEqual(recalculated["Portfolio"].cell(17, 2).value, "inside")
            self.assertTrue(recalculated["Portfolio"].cell(49, 2).value)
            self.assertEqual(recalculated["Portfolio"].cell(17, 3).value, "inside")

    def _returns(self):
        scenarios = build_scenarios(self.curves)
        beta = self.curves["2023-07-31"]
        returns = {}
        for scenario in SCENARIO_ORDER:
            returns[scenario] = {}
            for name, instrument in self.instruments.items():
                per_100 = full_reval_pnl_per_100(instrument, beta, scenarios[scenario]["dz"])
                returns[scenario][name] = scenario_return(instrument, per_100)
        return returns

    def test_cost_cushion_comparison(self):
        """Higher one-sided cost, and a larger comparison trade, reuse the same books."""
        current = normalize_weights(solve_two_bond(self.instruments, "S06M", "S30Y"), INSTRUMENT_IDS)
        benchmark = normalize_weights(solve_benchmark(self.instruments), INSTRUMENT_IDS)
        returns = self._returns()
        benchmark_pnl = {scenario: book_pnl(benchmark, returns[scenario]) for scenario in SCENARIO_ORDER}
        screen = clearing_analysis(current, benchmark, self.instruments, returns, "S05Y")
        shown = screen["turnover_just"] + 0.01
        specs = (
            (shown, 2.0 / 10_000.0, -981_168.373210, 18_831.626790, 76.281108),
            (shown, 4.2 / 10_000.0, -1_000_772.275230, -772.275230, 76.281108),
            (0.50, 2.0 / 10_000.0, -878_617.412038, 121_382.587962, 71.971970),
            (0.50, 4.2 / 10_000.0, -900_617.412038, 99_382.587962, 71.971970),
        )
        rows = []
        for turnover, cost_rate, worst, cushion, convexity in specs:
            row = candidate_stress_cushion(
                current,
                self.instruments,
                returns,
                benchmark_pnl,
                "S05Y",
                turnover,
                cost_rate,
                SCENARIO_ORDER,
            )
            self.assertTrue(row["long_only"])
            self.assertAlmostEqual(row["dv01"], DV01_TARGET, delta=1e-6)
            self.assertEqual(row["worst_scenario"], "hist_2023_03_08_to_2023_03_13")
            self.assertAlmostEqual(row["worst_relative_pnl"], worst, delta=0.02)
            self.assertAlmostEqual(row["cushion_vs_budget"], cushion, delta=0.02)
            self.assertAlmostEqual(row["convexity"], convexity, delta=1e-4)
            rows.append(row)
        self.assertAlmostEqual(rows[0]["cushion_vs_budget"], self.case["expanded_s05y"]["cushion_dollars"], delta=1e-4)
        self.assertLess(rows[1]["cushion_vs_budget"], 0.0)
        self.assertGreater(rows[3]["cushion_vs_budget"], 0.0)
        self.assertLess(rows[2]["convexity"], rows[0]["convexity"])
        self.assertNotAlmostEqual(rows[3]["turnover"], rows[1]["turnover"], places=4)

    def test_budget_changes_the_lowest_turnover_buy(self):
        """Same six-name search at three relative-loss budgets. Not a new optimizer."""
        current = normalize_weights(solve_two_bond(self.instruments, "S06M", "S30Y"), INSTRUMENT_IDS)
        benchmark = normalize_weights(solve_benchmark(self.instruments), INSTRUMENT_IDS)
        returns = self._returns()
        expected = (
            (500_000.0, "S07Y", 0.6693376628, "hist_2023_03_08_to_2023_03_13"),
            (1_000_000.0, "S05Y", 0.4355432277, "hist_2023_03_08_to_2023_03_13"),
            (1_500_000.0, "S05Y", 0.1752569620, "twist_steepener"),
        )
        for budget, buy, turnover, binding in expected:
            found = lowest_clearing_buy(current, benchmark, self.instruments, returns, budget=budget)
            self.assertTrue(found["feasible"])
            self.assertEqual(found["buy"], buy)
            self.assertEqual(found["binding_lower_scenario"], binding)
            self.assertAlmostEqual(found["turnover_just"], turnover, places=8)
        base = lowest_clearing_buy(current, benchmark, self.instruments, returns)
        self.assertEqual(base["buy"], self.case["case_decision"]["lowest_feasible_direction"]["buy"])
        self.assertAlmostEqual(
            base["turnover_just"],
            self.case["expanded_s05y"]["turnover_just"],
            places=8,
        )

    def test_front_bucket_kr01_fails_node_comparison(self):
        import validate_independently as independent

        with (ROOT / "outputs" / "instruments.csv").open(newline="") as handle:
            published = {row["instrument"]: row for row in csv.DictReader(handle)}
        bad = {}
        for name, row in published.items():
            parallel = float(row["dv01_per_100"])
            bad[name] = [parallel, 0.0, 0.0, 0.0, 0.0, 0.0]
            self.assertAlmostEqual(sum(bad[name]), parallel, places=12)
        self.assertFalse(independent.bond_kr01_within_tolerance(bad, published))
        self.assertGreater(independent.bond_kr01_max_abs_gap(bad, published), 0.1)
        curves = independent.load_curves()
        beta = curves["2023-07-31"]
        good = {}
        for name, coupon, maturity in independent.CONTRACTS:
            good[name] = independent.kr01(independent.cashflows(coupon, maturity), beta)
        self.assertTrue(independent.bond_kr01_within_tolerance(good, published))


if __name__ == "__main__":
    unittest.main()
